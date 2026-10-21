"""Export the saved cost model to budget_values.json for the deck builder.

Run after editing and SAVING the spreadsheet in Excel (Excel stores the computed
values): python docs/budget/export_budget.py
"""
import json
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).parent
book = load_workbook(HERE / 'MachineCare_Cost_Comparison.xlsx', data_only=True)
inputs = {row[0]: row[1] for row in book['Inputs'].iter_rows(min_row=5, values_only=True) if row[0] and row[1] is not None}
lines, side = {'vendor': {}, 'machinecare': {}, 'cumulative': {}, 'per_user_month': {}, 'reconciliation': {}}, None
for row in book['Comparison'].iter_rows(min_row=5, values_only=True):
    label = row[0]
    if not label:
        continue
    if label in ('OUTSIDE VENDOR', 'MACHINECARE (IN-HOUSE)'):
        side = 'vendor' if label == 'OUTSIDE VENDOR' else 'machinecare'
    elif side and isinstance(row[2], (int, float)):
        lines[side][label] = {'basis': row[1], 'years': list(row[2:7]), 'total': row[7]}
        if label.endswith('total per year'):
            side = None
    elif label.startswith('Cumulative'):
        lines['cumulative'][label] = list(row[2:7])
    elif label in ('Outside vendor', 'MachineCare'):
        lines['per_user_month'][label] = row[2]
    elif label.endswith('without added lines (security, re-checks, admin)') or label.endswith('without added line (yearly re-check)'):
        lines['reconciliation'][label] = row[7]
categories = [dict(zip(['category', 'vendor_y1', 'vendor_later', 'vendor_5y', 'mc_y1', 'mc_later', 'mc_5y', 'difference'], row))
              for row in book['Budget detail'].iter_rows(min_row=5, values_only=True) if row[0] and isinstance(row[1], (int, float))]
missing = [r for r in categories if any(v is None for v in r.values())]
if not categories or missing:
    raise SystemExit('No computed values found. Open the workbook in Excel, save it, then run this again.')
(HERE / 'budget_values.json').write_text(json.dumps({'inputs': inputs, 'lines': lines, 'categories': categories}, indent=1, default=str))
print('Wrote', HERE / 'budget_values.json')
